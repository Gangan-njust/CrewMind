"""Word 表格样式工具"""
from docx.document import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


def _border_spec(val: str, sz: str = "6", color: str = "000000") -> dict[str, str]:
  return {"val": val, "sz": sz, "color": color, "space": "0"}


_BORDER_NONE = _border_spec("nil", "0")
_BORDER_THIN = _border_spec("single", "6")
_BORDER_THICK = _border_spec("single", "12")


def _set_cell_borders(cell, *, top, bottom, left, right) -> None:
  tc = cell._tc
  tc_pr = tc.get_or_add_tcPr()
  existing = tc_pr.find(qn("w:tcBorders"))
  if existing is not None:
    tc_pr.remove(existing)
  borders = OxmlElement("w:tcBorders")
  for edge, spec in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
    element = OxmlElement(f"w:{edge}")
    for key, value in spec.items():
      element.set(qn(f"w:{key}"), value)
    borders.append(element)
  tc_pr.append(borders)


def apply_three_line_table(table: Table) -> None:
  """将 Word 表格设为学术三线表：顶线、表头下线、底线，无竖线与行间线。"""
  rows = table.rows
  if not rows:
    return

  row_count = len(rows)
  for row_index, row in enumerate(rows):
    if row_index == 0:
      top, bottom = _BORDER_THICK, _BORDER_THIN if row_count > 1 else _BORDER_THICK
    elif row_index == row_count - 1:
      top, bottom = _BORDER_NONE, _BORDER_THICK
    else:
      top, bottom = _BORDER_NONE, _BORDER_NONE

    for cell in row.cells:
      _set_cell_borders(cell, top=top, bottom=bottom, left=_BORDER_NONE, right=_BORDER_NONE)
      for paragraph in cell.paragraphs:
        _normalize_table_paragraph(paragraph, is_header=row_index == 0)


def apply_three_line_tables(document: Document) -> None:
  for table in document.tables:
    apply_three_line_table(table)


def apply_black_fonts(document: Document) -> None:
  """将文档内全部文字设为黑色，避免导出后出现彩色或主题色字体。"""
  black = RGBColor(0, 0, 0)

  for style in document.styles:
    try:
      style.font.color.rgb = black
    except (AttributeError, KeyError, TypeError):
      pass

  def _colorize(paragraphs) -> None:
    for paragraph in paragraphs:
      for run in paragraph.runs:
        run.font.color.rgb = black

  _colorize(document.paragraphs)
  for table in document.tables:
    for row in table.rows:
      for cell in row.cells:
        _colorize(cell.paragraphs)
  for section in document.sections:
    for part in (section.header, section.footer):
      if part is not None:
        _colorize(part.paragraphs)


def _normalize_table_paragraph(paragraph: Paragraph, *, is_header: bool) -> None:
  for run in paragraph.runs:
    if is_header and run.bold is None:
      run.bold = True
