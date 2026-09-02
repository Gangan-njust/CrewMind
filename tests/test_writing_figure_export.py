"""图表占位导出清理与段落缩进跳过单测。"""
import io

from docx import Document

from backend.writing.export import build_markdown, build_docx
from backend.writing.paragraph_format import ensure_paragraph_first_indent


def _project(content: str) -> dict:
  return {
    "title": "示例论文",
    "paper_type": "journal",
    "updated_at": "2026-01-01T00:00:00",
    "sections": [
      {
        "id": "s1",
        "section_type": "results",
        "title": "结果",
        "sort_order": 0,
        "content": content,
      },
    ],
  }


def test_build_markdown_strips_figure_placeholders():
  content = (
    "实验结果表明准确率提升 12.3%。\n\n"
    "<!-- figure:results:2: 建议在此插入柱状图 -->\n"
    "![图：对比结果（待插入）](待插入图表)\n\n"
    "![真实插图](/figs/real.png)"
  )
  md = build_markdown(_project(content))
  assert "<!-- figure" not in md
  assert "待插入图表" not in md
  assert "/figs/real.png" in md  # 真实图片保留


def test_build_markdown_keep_image_placeholders():
  content = (
    "正文段落。\n\n"
    "<!-- figure:results:2: 建议在此插入柱状图 -->\n"
    "![图：对比结果（待插入）](待插入图表)"
  )
  md = build_markdown(_project(content), keep_image_placeholders=True)
  assert "<!-- figure" not in md  # 注释始终剔除
  assert "待插入图表" in md  # 图片占位按需保留


def test_build_markdown_strips_comment_inside_paragraph():
  content = "定量对比结果如下：<!-- figure: 建议插柱状图 --> 准确率 95.2%。"
  md = build_markdown(_project(content))
  assert "figure" not in md
  assert "95.2%" in md


def test_paragraph_indent_skips_comment_and_image_blocks():
  content = (
    "普通段落首行会被缩进。\n\n"
    "<!-- figure:methods:1: 建议插入流程图 -->\n\n"
    "![图注](待插入图表)\n\n"
    "| 方法 | 指标 |\n| --- | --- |"
  )
  cleaned = ensure_paragraph_first_indent(content)
  assert cleaned.split("\n")[0].startswith("　　")  # 普通段落被缩进
  for marker in ("<!-- figure", "![图注]", "| 方法 |"):
    for line in cleaned.split("\n"):
      if marker in line:
        assert not line.startswith("　　"), f"{marker} 不应被缩进"


def test_build_docx_no_broken_image_marker():
  content = (
    "实验结果表明准确率提升 12.3%。\n\n"
    "<!-- figure:results:2: 建议在此插入柱状图 -->\n"
    "![图：对比结果（待插入）](待插入图表)"
  )
  docx_bytes = build_docx(_project(content))
  doc = Document(io.BytesIO(docx_bytes))
  text = "\n".join(p.text for p in doc.paragraphs)
  assert "<image" not in text
  assert "待插入图表" not in text


def test_build_docx_embeds_cmasset_real_image(tmp_path, monkeypatch):
  """cmasset:// 真实素材在导出 docx 时内嵌为图片（不再出现引用文本）。"""
  import base64

  from backend.storage.writing_store import writing_store

  png_bytes = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
  )
  png = tmp_path / "real-figure.png"
  png.write_bytes(png_bytes)
  monkeypatch.setattr(
    writing_store, "resolve_asset_path",
    lambda asset_id: png if asset_id == "asset-1" else None,
  )

  project = _project("定量结果如图 1 所示。\n\n![图1：实验对比](cmasset://asset-1)")
  project["id"] = "project-x"
  docx_bytes = build_docx(project)
  doc = Document(io.BytesIO(docx_bytes))
  assert len(doc.inline_shapes) == 1
  text = "\n".join(p.text for p in doc.paragraphs)
  assert "cmasset" not in text


def test_build_docx_drops_missing_cmasset(monkeypatch):
  """cmasset 文件丢失时导出不报错，也不产生坏图标记。"""
  from backend.storage.writing_store import writing_store

  monkeypatch.setattr(writing_store, "resolve_asset_path", lambda asset_id: None)
  project = _project("定量结果如图 1 所示。\n\n![图1：实验对比](cmasset://asset-gone)")
  project["id"] = "project-x"
  docx_bytes = build_docx(project)
  doc = Document(io.BytesIO(docx_bytes))
  text = "\n".join(p.text for p in doc.paragraphs)
  assert "cmasset" not in text
  assert "<image" not in text

