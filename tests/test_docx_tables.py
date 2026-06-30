from docx import Document

from backend.utils.docx_tables import apply_three_line_table, apply_three_line_tables


def test_apply_three_line_table_sets_header_and_body_borders():
  document = Document()
  table = document.add_table(rows=3, cols=2)
  table.cell(0, 0).text = "指标"
  table.cell(0, 1).text = "数值"
  table.cell(1, 0).text = "准确率"
  table.cell(1, 1).text = "95%"
  table.cell(2, 0).text = "召回率"
  table.cell(2, 1).text = "92%"

  apply_three_line_table(table)

  header_top = table.cell(0, 0)._tc.tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders")
  assert header_top is not None
  borders = {child.tag.split("}")[-1]: child.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
             for child in header_top}
  assert borders["top"] == "single"
  assert borders["bottom"] == "single"
  assert borders["left"] == "nil"
  assert borders["right"] == "nil"

  body_top = table.cell(1, 0)._tc.tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders")
  body_vals = {child.tag.split("}")[-1]: child.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
               for child in body_top}
  assert body_vals["top"] == "nil"
  assert body_vals["bottom"] == "nil"

  last_bottom = table.cell(2, 0)._tc.tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders")
  last_vals = {child.tag.split("}")[-1]: child.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
               for child in last_bottom}
  assert last_vals["bottom"] == "single"


def test_apply_three_line_tables_on_document():
  document = Document()
  document.add_table(rows=1, cols=1)
  document.add_table(rows=2, cols=1)
  apply_three_line_tables(document)
  assert len(document.tables) == 2
